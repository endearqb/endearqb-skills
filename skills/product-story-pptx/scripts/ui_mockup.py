#!/usr/bin/env python3
"""Editable desktop-app UI mockup primitives for product-story-pptx.

The renderer intentionally builds native PowerPoint shapes.  It is not a
screenshot generator and must remain labelled as an interface mockup.
"""
from __future__ import annotations

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt


def _rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value.lstrip("#").upper())


def _name(shape, role: str, instance: str) -> None:
    shape.name = f"PST_UI_MOCK_{role}::{instance}"


def _remove_theme_effect(shape) -> None:
    shape.shadow.inherit = False
    for node in shape._element.xpath("./p:style/a:effectRef"):
        node.set("idx", "0")


def _set_cjk(run, font_cjk: str) -> None:
    rpr = run._r.get_or_add_rPr()
    ea = rpr.find(qn("a:ea"))
    if ea is None:
        ea = rpr.makeelement(qn("a:ea"), {})
        rpr.append(ea)
    ea.set("typeface", font_cjk)


def _shape(slide, rect, fill: str, line: str | None, role: str, instance: str,
           corner: float = 0.0):
    x, y, w, h = rect
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if corner else MSO_SHAPE.RECTANGLE
    shp = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    _remove_theme_effect(shp)
    shp.fill.solid()
    shp.fill.fore_color.rgb = _rgb(fill)
    if line:
        shp.line.color.rgb = _rgb(line)
        shp.line.width = Pt(1)
    else:
        shp.line.fill.background()
    if corner:
        short = min(w, h)
        shp.adjustments[0] = min(0.12, corner / short) if short else 0.04
    _name(shp, role, instance)
    return shp


def _text(slide, rect, text: str, size: float, color: str, fonts: dict,
          role: str, instance: str, *, bold=False, align=PP_ALIGN.LEFT):
    x, y, w, h = rect
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    _name(box, role, instance)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.03)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_before = p.space_after = Pt(0)
    run = p.add_run()
    run.text = str(text)
    run.font.name = fonts["font_latin"]
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = _rgb(color)
    _set_cjk(run, fonts["font_cjk"])
    return box


def _line(slide, x1, y1, x2, y2, color: str, role: str, instance: str):
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    line.line.color.rgb = _rgb(color)
    line.line.width = Pt(0.8)
    _name(line, role, instance)
    return line


def _colors(theme: str, palette: dict) -> dict:
    if theme == "dark":
        return {
            "shell": "#132026", "topbar": "#203138", "surface": "#203138",
            "surface2": "#2A3C43", "text": "#FFFFFF", "muted": "#78C7B8",
            "accent": palette["accent"], "rule": "#455860",
        }
    return {
        "shell": palette["background"], "topbar": palette["surface"],
        "surface": palette["surface"], "surface2": "#FFFFFF",
        "text": palette["text"], "muted": palette["text_muted"],
        "accent": palette["accent"], "rule": palette["rule"],
    }

def draw_ui_mockup(slide, rect, spec: dict, palette: dict, fonts: dict,
                   instance: str) -> dict:
    """Draw one editable desktop-app mockup and return its geometry metadata."""
    x, y, w, h = rect
    theme = spec.get("theme", "light")
    c = _colors(theme, palette)
    corner = float(spec.get("corner_in", 0.18))
    corner = max(0.10, min(0.25, corner))
    titlebar_h = 0.50
    inset = 0.02

    shell = _shape(slide, rect, c["shell"], c["rule"], "SHELL", instance, corner)
    topbar = _shape(
        slide, (x + inset, y + inset, w - 2 * inset, titlebar_h - inset),
        c["topbar"], None, "TOPBAR", instance, 0.06,
    )
    _shape(
        slide, (x + inset, y + 0.27, w - 2 * inset, titlebar_h - 0.27),
        c["topbar"], None, "TOPBAR_JOIN", instance,
    )
    _line(slide, x + inset, y + titlebar_h, x + w - inset, y + titlebar_h,
          c["rule"], "DIVIDER", instance)

    for idx, color in enumerate(("#E6605C", "#F3B74E", "#00A676")):
        dot = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(x + 0.17 + idx * 0.18), Inches(y + 0.18),
            Inches(0.075), Inches(0.075),
        )
        dot.fill.solid()
        dot.fill.fore_color.rgb = _rgb(color)
        dot.line.fill.background()
        _remove_theme_effect(dot)
        _name(dot, f"DOT_{idx + 1}", instance)

    _text(slide, (x + 0.70, y + 0.08, w - 1.85, 0.30),
          spec.get("title", "Product App"), 12, c["text"], fonts,
          "TITLE", instance, bold=True)
    badge_text = spec.get("badge", "界面示意")
    _shape(slide, (x + w - 1.03, y + 0.08, 0.84, 0.30),
           c["surface2"], c["rule"], "BADGE", instance, 0.06)
    _text(slide, (x + w - 1.01, y + 0.08, 0.80, 0.30), badge_text, 9,
          c["muted"], fonts, "BADGE_TEXT", instance, bold=True,
          align=PP_ALIGN.CENTER)

    cx, cy = x + 0.18, y + 0.64
    cw, ch = w - 0.36, h - 0.82
    sidebar = spec.get("sidebar") or []
    if sidebar:
        sw = min(1.30, max(1.05, cw * 0.22))
        _shape(slide, (cx, cy, sw, ch), c["surface"], None,
               "SIDEBAR", instance, 0.10)
        for i, label in enumerate(sidebar[:5]):
            _text(slide, (cx + 0.10, cy + 0.20 + i * 0.52, sw - 0.20, 0.28),
                  label, 9, c["accent"] if i == 0 else c["muted"], fonts,
                  f"SIDEBAR_ITEM_{i + 1}", instance, bold=i == 0)
        cx += sw + 0.22
        cw -= sw + 0.22

    headline = spec.get("headline")
    subhead = spec.get("subhead")
    if headline:
        _text(slide, (cx, cy, cw, 0.34), headline, 13, c["text"], fonts,
              "HEADLINE", instance, bold=True)
        cy += 0.36
        ch -= 0.36
    if subhead:
        _text(slide, (cx, cy, cw, 0.30), subhead, 9, c["muted"], fonts,
              "SUBHEAD", instance)
        cy += 0.36
        ch -= 0.36

    kpis = (spec.get("kpis") or [])[:3]
    if kpis:
        gap = 0.16
        kw = (cw - gap * (len(kpis) - 1)) / len(kpis)
        for i, item in enumerate(kpis):
            xx = cx + i * (kw + gap)
            highlighted = bool(item.get("highlight"))
            _shape(slide, (xx, cy, kw, 0.78), c["surface2"],
                   c["accent"] if highlighted else c["rule"],
                   f"KPI_{i + 1}", instance, 0.09)
            _text(slide, (xx + 0.10, cy + 0.08, kw - 0.20, 0.23),
                  item.get("label", "指标"), 8.5, c["muted"], fonts,
                  f"KPI_LABEL_{i + 1}", instance)
            _text(slide, (xx + 0.10, cy + 0.34, kw - 0.20, 0.30),
                  item.get("value", "—"), 14, c["text"], fonts,
                  f"KPI_VALUE_{i + 1}", instance, bold=True)
        cy += 0.96
        ch -= 0.96

    items = (spec.get("items") or [])[:4]
    if items:
        row_gap = 0.12
        row_h = min(0.72, max(0.48, (ch - row_gap * (len(items) - 1)) / len(items)))
        for i, item in enumerate(items):
            yy = cy + i * (row_h + row_gap)
            highlighted = bool(item.get("highlight"))
            fill = c["accent"] if highlighted else c["surface"]
            text_color = "#FFFFFF" if highlighted else c["text"]
            _shape(slide, (cx, yy, cw, row_h), fill,
                   c["accent"] if highlighted else c["rule"],
                   f"ROW_{i + 1}", instance, 0.08)
            _text(slide, (cx + 0.14, yy + 0.06, cw * 0.58, row_h - 0.12),
                  item.get("label", f"Item {i + 1}"), 10, text_color, fonts,
                  f"ROW_LABEL_{i + 1}", instance, bold=True)
            value = item.get("value") or item.get("status") or ""
            if value:
                _text(slide, (cx + cw * 0.62, yy + 0.06, cw * 0.34, row_h - 0.12),
                      value, 9, "#FFFFFF" if highlighted else c["accent"], fonts,
                      f"ROW_VALUE_{i + 1}", instance, bold=True,
                      align=PP_ALIGN.RIGHT)

    return {
        "instance": instance,
        "theme": theme,
        "corner_in": corner,
        "titlebar_h": titlebar_h,
        "shell_name": shell.name,
    }
