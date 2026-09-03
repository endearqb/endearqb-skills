#!/usr/bin/env python3
"""Draw restrained, editable flow diagrams with native PowerPoint shapes.

The module intentionally avoids connector objects. Connectors often re-route or
shift between PowerPoint and LibreOffice; small chevrons and deterministic
coordinates are more portable for generated decks.

Public API:
    draw_flow(slide, rect, steps, palette, fonts, sizes, geo,
              direction="horizontal", style="cards")

Supported styles:
    cards      spaced rounded cards with numbered badges
    chevron    contiguous directional band
    vertical   vertical timeline
    layers     stacked architecture bands
"""
from __future__ import annotations

from collections.abc import Mapping, Sequence
from typing import Any

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

Rect = tuple[float, float, float, float]

_ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
_ANCHOR = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}


def _rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value.lstrip("#").upper())


def _set_cjk_font(run: Any, font_cjk: str) -> None:
    r_pr = run._r.get_or_add_rPr()
    ea = r_pr.find(qn("a:ea"))
    if ea is None:
        ea = r_pr.makeelement(qn("a:ea"), {})
        r_pr.append(ea)
    ea.set("typeface", font_cjk)


def _shape(
    slide: Any,
    kind: Any,
    rect: Rect,
    *,
    fill: str | None = None,
    line: str | None = None,
    line_w: float = 0.75,
    corner: float | None = None,
    name: str | None = None,
) -> Any:
    x, y, w, h = rect
    if w <= 0 or h <= 0:
        raise ValueError(f"invalid shape rect: {rect}")
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        shape.name = name
    shape.shadow.inherit = False
    if fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = _rgb(fill)
    else:
        shape.fill.background()
    if line:
        shape.line.color.rgb = _rgb(line)
        shape.line.width = Pt(line_w)
    else:
        shape.line.fill.background()
    if corner is not None and kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        short = min(w, h)
        shape.adjustments[0] = max(0.0, min(0.5, corner / short)) if short else 0.1
    return shape


def _text(
    slide: Any,
    rect: Rect | None,
    text: str,
    size: float,
    color: str,
    fonts: Mapping[str, str],
    *,
    bold: bool = False,
    align: str = "left",
    anchor: str = "top",
    margin: float = 0.05,
    word_wrap: bool = True,
    shape: Any | None = None,
    name: str | None = None,
) -> Any:
    """Add text to a new text box or an existing auto-shape.

    The original draft unpacked ``rect`` even when ``shape`` was supplied,
    which made every numbered badge fail with ``TypeError: cannot unpack None``.
    """
    if shape is None:
        if rect is None:
            raise ValueError("rect is required when shape is not supplied")
        x, y, w, h = rect
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        shape.name = name
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = word_wrap
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = _ANCHOR[anchor]
    p = tf.paragraphs[0]
    p.alignment = _ALIGN[align]
    run = p.add_run()
    run.text = str(text)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = fonts.get("font_latin", "Arial")
    run.font.color.rgb = _rgb(color)
    _set_cjk_font(run, fonts.get("font_cjk", "Microsoft YaHei"))
    return shape


def _single_line_size(text: str, base_size: float, width_in: float, *, minimum: float = 10.0) -> float:
    """Conservatively fit a short label on one line without auto-wrap.

    CJK glyphs are roughly square at presentation sizes; Latin and punctuation
    consume less width. The estimate intentionally leaves a small safety margin
    for LibreOffice/PowerPoint font fallback differences.
    """
    units = sum(1.0 if ord(ch) > 127 else 0.58 for ch in str(text) if not ch.isspace())
    if units <= 0 or width_in <= 0:
        return base_size
    available_pt = width_in * 72.0 * 0.88
    return max(minimum, min(base_size, available_pt / units))


def _badge(
    slide: Any,
    cx: float,
    cy: float,
    diameter: float,
    number: int,
    *,
    fill: str,
    text_color: str,
    fonts: Mapping[str, str],
    size: float,
    line: str | None = None,
) -> Any:
    shape = _shape(
        slide,
        MSO_SHAPE.OVAL,
        (cx - diameter / 2, cy - diameter / 2, diameter, diameter),
        fill=fill,
        line=line,
        name=f"PST_FLOW_BADGE_{number}",
    )
    _text(
        slide,
        None,
        str(number),
        size,
        text_color,
        fonts,
        bold=True,
        align="center",
        anchor="middle",
        margin=0.0,
        shape=shape,
    )
    return shape


def _arrow(slide: Any, rect: Rect, color: str, number: int) -> Any:
    return _shape(
        slide,
        MSO_SHAPE.CHEVRON,
        rect,
        fill=color,
        name=f"PST_FLOW_ARROW_{number}",
    )


def _cards(slide: Any, rect: Rect, steps: Sequence[Mapping[str, Any]], pal: Mapping[str, str],
           fonts: Mapping[str, str], sizes: Mapping[str, float], g: Mapping[str, float]) -> None:
    x, y, w, h = rect
    n = len(steps)
    # Preserve breathing room for 3–5 nodes, then tighten progressively for 6–7.
    compression = max(0, n - 5)
    gap = max(0.12, float(g["gap"]) - 0.07 * compression)
    arrow_w = max(0.20, float(g["arrow_w"]) - 0.04 * compression)
    node_w = (w - (n - 1) * (2 * gap + arrow_w)) / n
    if node_w < 0.95:
        raise ValueError(
            f"flow cards are too narrow ({node_w:.2f} in for {n} nodes); split the flow or use layers"
        )
    node_h = min(h, float(g["node_h_max"]))
    top = y + (h - node_h) / 2
    badge_d = min(float(g["badge_d"]), node_w * 0.28)
    label_size = max(12, float(sizes["flow_label"]) - compression)
    desc_size = max(10, float(sizes["flow_desc"]) - 0.5 * compression)
    cursor_x = x
    for index, step in enumerate(steps, start=1):
        highlighted = bool(step.get("highlight"))
        fill = pal["hl_fill"] if highlighted else pal["node_fill"]
        line = None if highlighted else pal["node_line"]
        text_color = pal["hl_text"] if highlighted else pal["text"]
        muted = pal["hl_text"] if highlighted else pal["muted"]
        _shape(
            slide,
            MSO_SHAPE.ROUNDED_RECTANGLE,
            (cursor_x, top, node_w, node_h),
            fill=fill,
            line=line,
            corner=float(g["corner"]),
            name=f"PST_FLOW_NODE_{index}",
        )
        badge_cx = cursor_x + 0.24 + badge_d / 2
        badge_cy = top + 0.24 + badge_d / 2
        _badge(
            slide,
            badge_cx,
            badge_cy,
            badge_d,
            index,
            fill=pal["hl_text"] if highlighted else pal["node_line"],
            text_color=pal["hl_fill"] if highlighted else pal["text"],
            fonts=fonts,
            size=float(sizes["flow_badge"]),
        )
        label_y = top + 0.24 + badge_d + 0.14
        inner_pad = 0.16 if n >= 6 else 0.24
        inner_w = node_w - 2 * inner_pad
        _text(
            slide,
            (cursor_x + inner_pad, label_y, inner_w, 0.58),
            str(step["label"]),
            _single_line_size(str(step["label"]), label_size, inner_w, minimum=12.0),
            text_color,
            fonts,
            bold=True,
            margin=0.0,
            word_wrap=False,
            name=f"PST_FLOW_LABEL_{index}",
        )
        desc = step.get("desc")
        if desc:
            _text(
                slide,
                (cursor_x + inner_pad, label_y + 0.62, inner_w, max(0.45, node_h - (label_y - top) - 0.72)),
                str(desc),
                _single_line_size(str(desc), desc_size, inner_w, minimum=12.0),
                muted,
                fonts,
                margin=0.0,
                word_wrap=False,
                name=f"PST_FLOW_DESC_{index}",
            )
        if index < n:
            arrow_x = cursor_x + node_w + gap
            arrow_h = 0.38
            _arrow(slide, (arrow_x, top + node_h / 2 - arrow_h / 2, arrow_w, arrow_h), pal["arrow"], index)
        cursor_x += node_w + 2 * gap + arrow_w


def _chevron(slide: Any, rect: Rect, steps: Sequence[Mapping[str, Any]], pal: Mapping[str, str],
             fonts: Mapping[str, str], sizes: Mapping[str, float], g: Mapping[str, float]) -> None:
    x, y, w, h = rect
    n = len(steps)
    node_h = min(float(g["node_h"]), h * 0.46)
    overlap = float(g["overlap"])
    node_w = (w + (n - 1) * overlap) / n
    if node_w < 1.05:
        raise ValueError("chevron flow is too dense; split the flow or use a vertical layout")
    top = y + max(0.2, (h - node_h - 1.15) / 2)
    cursor_x = x
    for index, step in enumerate(steps, start=1):
        highlighted = bool(step.get("highlight"))
        kind = MSO_SHAPE.PENTAGON if index == 1 else MSO_SHAPE.CHEVRON
        fill = pal["hl_fill"] if highlighted else pal["node_fill"]
        shape = _shape(
            slide,
            kind,
            (cursor_x, top, node_w, node_h),
            fill=fill,
            line=None if highlighted else pal["node_line"],
            name=f"PST_FLOW_CHEVRON_{index}",
        )
        text_color = pal["hl_text"] if highlighted else pal["text"]
        number_w = 0.30
        left_pad = 0.22 if index == 1 else 0.30
        right_pad = 0.30
        label_x = cursor_x + left_pad + number_w
        label_w = node_w - left_pad - right_pad - number_w
        _text(
            slide,
            (cursor_x + left_pad, top, number_w, node_h),
            str(index),
            float(sizes["flow_badge"]),
            text_color,
            fonts,
            bold=True,
            align="center",
            anchor="middle",
            margin=0.0,
            word_wrap=False,
            name=f"PST_FLOW_INDEX_{index}",
        )
        _text(
            slide,
            (label_x, top, label_w, node_h),
            str(step["label"]),
            _single_line_size(str(step["label"]), float(sizes["flow_label"]), label_w, minimum=12.0),
            text_color,
            fonts,
            bold=highlighted,
            align="center",
            anchor="middle",
            margin=0.0,
            word_wrap=False,
            name=f"PST_FLOW_LABEL_{index}",
        )
        if step.get("desc"):
            _text(
                slide,
                (cursor_x + 0.15, top + node_h + 0.18, node_w - 0.36, 0.84),
                str(step["desc"]),
                _single_line_size(str(step["desc"]), float(sizes["flow_desc"]), node_w - 0.36, minimum=12.0),
                pal["hl_fill"] if highlighted else pal["muted"],
                fonts,
                align="center",
                margin=0.0,
                word_wrap=False,
                name=f"PST_FLOW_DESC_{index}",
            )
        cursor_x += node_w - overlap


def _vertical(slide: Any, rect: Rect, steps: Sequence[Mapping[str, Any]], pal: Mapping[str, str],
              fonts: Mapping[str, str], sizes: Mapping[str, float], g: Mapping[str, float]) -> None:
    x, y, w, h = rect
    n = len(steps)
    row_h = min(h / n, float(g["row_h_max"]))
    total = row_h * n
    top = y + (h - total) / 2
    line_x = x + float(g["line_x_offset"])
    badge_d = float(g["badge_d"])
    if n > 1:
        _shape(
            slide,
            MSO_SHAPE.RECTANGLE,
            (line_x - 0.012, top + row_h / 2, 0.024, total - row_h),
            fill=pal["node_line"],
            name="PST_FLOW_TIMELINE",
        )
    text_x = line_x + badge_d / 2 + 0.30
    text_w = w - (text_x - x)
    for index, step in enumerate(steps, start=1):
        highlighted = bool(step.get("highlight"))
        center_y = top + row_h * (index - 0.5)
        _badge(
            slide,
            line_x,
            center_y,
            badge_d,
            index,
            fill=pal["hl_fill"] if highlighted else pal["node_fill"],
            text_color=pal["hl_text"] if highlighted else pal["text"],
            line=None if highlighted else pal["node_line"],
            fonts=fonts,
            size=float(sizes["flow_badge"]),
        )
        _text(
            slide,
            (text_x, center_y - row_h / 2 + 0.03, text_w, row_h * 0.52),
            str(step["label"]),
            float(sizes["flow_label"]),
            pal["hl_fill"] if highlighted else pal["text"],
            fonts,
            bold=True,
            anchor="bottom",
            name=f"PST_FLOW_LABEL_{index}",
        )
        if step.get("desc"):
            _text(
                slide,
                (text_x, center_y + 0.02, text_w, max(0.30, row_h * 0.45)),
                str(step["desc"]),
                float(sizes["flow_desc"]),
                pal["muted"],
                fonts,
                name=f"PST_FLOW_DESC_{index}",
            )


def _layers(slide: Any, rect: Rect, steps: Sequence[Mapping[str, Any]], pal: Mapping[str, str],
            fonts: Mapping[str, str], sizes: Mapping[str, float], g: Mapping[str, float]) -> None:
    x, y, w, h = rect
    n = len(steps)
    gap = float(g["band_gap"])
    band_h = min((h - gap * (n - 1)) / n, float(g["band_h_max"]))
    if band_h < 0.55:
        raise ValueError("layer bands are too short; split the architecture across slides")
    total = band_h * n + gap * (n - 1)
    top = y + (h - total) / 2
    for index, step in enumerate(steps, start=1):
        highlighted = bool(step.get("highlight"))
        band_y = top + (index - 1) * (band_h + gap)
        _shape(
            slide,
            MSO_SHAPE.ROUNDED_RECTANGLE,
            (x, band_y, w, band_h),
            fill=pal["hl_fill"] if highlighted else pal["node_fill"],
            line=None if highlighted else pal["node_line"],
            corner=0.08,
            name=f"PST_FLOW_LAYER_{index}",
        )
        _text(
            slide,
            (x + 0.30, band_y, w * 0.31, band_h),
            str(step["label"]),
            float(sizes["flow_label"]),
            pal["hl_text"] if highlighted else pal["text"],
            fonts,
            bold=True,
            anchor="middle",
            name=f"PST_FLOW_LABEL_{index}",
        )
        if step.get("desc"):
            _text(
                slide,
                (x + w * 0.36, band_y, w * 0.60, band_h),
                str(step["desc"]),
                float(sizes["flow_desc"]),
                pal["hl_text"] if highlighted else pal["muted"],
                fonts,
                anchor="middle",
                name=f"PST_FLOW_DESC_{index}",
            )


_STYLES = {"cards": _cards, "chevron": _chevron, "vertical": _vertical, "layers": _layers}


def draw_flow(
    slide: Any,
    rect: Rect,
    steps: Sequence[Mapping[str, Any]],
    palette: Mapping[str, str],
    fonts: Mapping[str, str],
    sizes: Mapping[str, float],
    geo: Mapping[str, Mapping[str, float]],
    *,
    direction: str = "horizontal",
    style: str = "cards",
) -> None:
    if not 3 <= len(steps) <= 7:
        raise ValueError(f"flow requires 3–7 steps, got {len(steps)}")
    if direction == "vertical":
        style = "vertical"
    if style not in _STYLES:
        raise ValueError(f"unknown flow style {style!r}; choose one of {sorted(_STYLES)}")
    _STYLES[style](slide, rect, steps, palette, fonts, sizes, geo[style])


if __name__ == "__main__":  # pragma: no cover
    import sys
    from pptx import Presentation

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    demo_steps = [
        {"label": "授权数据源", "desc": "文档 / 代码 / 消息"},
        {"label": "建立索引", "desc": "自动解析与权限继承", "highlight": True},
        {"label": "自然提问", "desc": "答案附来源"},
    ]
    palette = {
        "node_fill": "#F3F5F7",
        "node_line": "#D7DBE0",
        "text": "#0B0C0F",
        "muted": "#6B7280",
        "hl_fill": "#00A676",
        "hl_text": "#FFFFFF",
        "arrow": "#C7CCD3",
    }
    geometry = {
        "cards": {"gap": 0.30, "arrow_w": 0.32, "node_h_max": 2.4, "badge_d": 0.42, "corner": 0.12},
        "chevron": {"node_h": 1.1, "overlap": 0.18},
        "vertical": {"line_x_offset": 0.3, "row_h_max": 1.15, "badge_d": 0.42},
        "layers": {"band_gap": 0.14, "band_h_max": 1.0},
    }
    chosen = sys.argv[1] if len(sys.argv) > 1 else "cards"
    draw_flow(
        slide,
        (0.8, 1.6, 11.733, 5.0),
        demo_steps,
        palette,
        {"font_latin": "Arial", "font_cjk": "Microsoft YaHei"},
        {"flow_label": 16, "flow_desc": 12, "flow_badge": 12},
        geometry,
        style=chosen,
    )
    output = f"flow_demo_{chosen}.pptx"
    prs.save(output)
    print(f"→ {output}")
