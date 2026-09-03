#!/usr/bin/env python3
"""build_deck.py — 由 deck spec 生成 .pptx。

用法:
    python scripts/build_deck.py spec.yaml out/deck.pptx
    python scripts/build_deck.py spec.yaml out/deck.pptx --no-lint     # 跳过 lint（调试用，不建议）
    python scripts/build_deck.py spec.yaml out/deck.pptx --keep-temp   # 保留裁切图 / 图表 PNG

退出码:
    0 成功；1 lint FAIL；2 spec / 资源错误

产物:
    out/deck.pptx
    out/deck.manifest.json    页序、type、bleed、energy、素材与 WARN 清单，用于 Step 7 交付说明

依赖: python-pptx pyyaml pillow matplotlib
"""
from __future__ import annotations

import argparse
import hashlib
import json
import math
import os
import re
import shutil
import sys
import tempfile

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from chart_swd import build_palette, render_chart  # noqa: E402
from flow_diagram import draw_flow  # noqa: E402
from lint_deck import CJK_RE, lint, load_registry, load_tokens, load_yaml  # noqa: E402
from ui_mockup import draw_ui_mockup  # noqa: E402

ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
ANCHOR = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}


# ---------------------------------------------------------------------------
# 颜色
# ---------------------------------------------------------------------------
def rgb(hex_str: str) -> RGBColor:
    return RGBColor.from_string(hex_str.lstrip("#").upper())


def hex_to_tuple(h: str) -> tuple[int, int, int]:
    h = h.lstrip("#")
    return int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)


def tuple_to_hex(t) -> str:
    return "#%02X%02X%02X" % tuple(max(0, min(255, int(round(c)))) for c in t)


def luminance(h: str) -> float:
    r, g, b = (c / 255 for c in hex_to_tuple(h))
    lin = [c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4 for c in (r, g, b)]
    return 0.2126 * lin[0] + 0.7152 * lin[1] + 0.0722 * lin[2]


def on_color(bg: str) -> str:
    """给一个背景色，返回可读的文字色（黑或白）。"""
    return "#111111" if luminance(bg) > 0.45 else "#FFFFFF"


def mix(a: str, b: str, t: float) -> str:
    ta, tb = hex_to_tuple(a), hex_to_tuple(b)
    return tuple_to_hex(tuple(ta[i] * (1 - t) + tb[i] * t for i in range(3)))


# ---------------------------------------------------------------------------
# 文本估算（无字体度量的保守近似；渲染后仍以 render_preview 为准）
# ---------------------------------------------------------------------------
def _char_units(text: str) -> float:
    u = 0.0
    for ch in text:
        if CJK_RE.match(ch):
            u += 1.0
        elif ch.isspace():
            u += 0.3
        elif ch.isupper() or ch.isdigit():
            u += 0.62
        else:
            u += 0.52
    return u


def estimate_lines(text: str, w_in: float, size_pt: float) -> int:
    em_in = size_pt / 72.0
    per_line = max(1.0, (w_in - 0.1) / em_in)
    lines = 0
    for para in str(text).split("\n"):
        lines += max(1, math.ceil(_char_units(para) / per_line))
    return lines


def estimate_height(text: str, w_in: float, size_pt: float, ls: float = 1.15, pad: float = 0.1) -> float:
    return estimate_lines(text, w_in, size_pt) * size_pt * ls / 72.0 + pad


def fit_size(text: str, w_in: float, h_in: float, size_pt: float, min_pt: float, ls: float) -> tuple[float, bool]:
    """缩字直到放得下或到下限。返回 (字号, 是否仍溢出)。"""
    s = float(size_pt)
    while estimate_height(text, w_in, s, ls) > h_in and s - 2 >= min_pt:
        s -= 2
    return s, estimate_height(text, w_in, s, ls) > h_in


# ---------------------------------------------------------------------------
# 图片
# ---------------------------------------------------------------------------
def cover_crop(src: str, w_in: float, h_in: float, focus: tuple[float, float], out_dir: str,
               max_w_px: int = 3200) -> str:
    """按目标宽高比裁切（不拉伸），以 focus 为锚点。返回临时文件路径。"""
    target = w_in / h_in
    with Image.open(src) as im:
        im.load()
        has_alpha = im.mode in ("RGBA", "LA") or (im.mode == "P" and "transparency" in im.info)
        im = im.convert("RGBA" if has_alpha else "RGB")
        W, H = im.size
        cur = W / H
        if cur > target:  # 太宽 → 裁左右
            cw, ch = int(round(H * target)), H
        else:  # 太高 → 裁上下
            cw, ch = W, int(round(W / target))
        fx, fy = focus
        x0 = int(round((W - cw) * fx))
        y0 = int(round((H - ch) * fy))
        im = im.crop((x0, y0, x0 + cw, y0 + ch))
        if im.size[0] > max_w_px:
            im = im.resize((max_w_px, int(round(max_w_px / target))), Image.LANCZOS)
        base = re.sub(r"[^A-Za-z0-9_.-]", "_", os.path.basename(src))
        out = os.path.join(out_dir, f"crop_{abs(hash((src, w_in, h_in, fx, fy)))}_{base}")
        out = os.path.splitext(out)[0] + (".png" if has_alpha else ".jpg")
        if has_alpha:
            im.save(out, "PNG", optimize=True)
        else:
            im.save(out, "JPEG", quality=90, optimize=True)
    return out


def gradient_png(out_path: str, color: str, alpha_from: float, alpha_to: float, cover_ratio: float,
                 size=(1333, 750)) -> str:
    """底部 → 顶部渐隐的遮罩 PNG（direction=bottom_to_top）。"""
    W, H = size
    r, g, b = hex_to_tuple(color)
    col = Image.new("RGBA", (1, H))
    px = col.load()
    start = H * (1 - cover_ratio)
    for y in range(H):
        if y < start:
            a = 0.0
        else:
            t = (y - start) / max(1.0, H - start)
            a = alpha_to + (alpha_from - alpha_to) * t
        px[0, y] = (r, g, b, int(round(a * 255)))
    col.resize((W, H), Image.NEAREST).save(out_path, "PNG")
    return out_path


# ---------------------------------------------------------------------------
# 构建上下文
# ---------------------------------------------------------------------------
class Ctx:
    def __init__(self, spec: dict, spec_path: str, reg: dict, tokens: dict, tmp: str):
        self.spec, self.reg, self.tokens, self.tmp = spec, reg, tokens, tmp
        deck = spec["deck"]
        self.deck = deck
        self.spec_dir = os.path.dirname(os.path.abspath(spec_path))
        self.assets_dir = os.path.join(self.spec_dir, deck.get("assets_dir", "./assets"))
        self.theme = dict(tokens["themes"][deck["theme"]])
        brand = deck.get("brand") or {}
        if brand.get("accent"):
            self.theme["accent"] = brand["accent"]
            self.theme["accent_soft"] = mix(brand["accent"], self.theme["background"], 0.45)
        if brand.get("color_block"):
            self.theme["color_block"] = brand["color_block"]
            self.theme["color_block_text"] = on_color(brand["color_block"])
        ty = tokens["typography"]
        self.fonts = {
            "font_latin": brand.get("font_latin") or ty["font_latin"],
            "font_cjk": brand.get("font_cjk") or ty["font_cjk"],
        }
        self.sizes = ty["sizes"]
        self.min_sizes = ty["min_sizes"]
        self.ls = ty.get("line_spacing", 1.15)
        self.canvas = (reg["canvas"]["width_in"], reg["canvas"]["height_in"])
        self.warnings: list[str] = []
        self.manifest_slides: list[dict] = []

    def warn(self, slide_no: int, msg: str) -> None:
        self.warnings.append(f"slide {slide_no}: {msg}")
        print(f"[WARN] slide {slide_no}: {msg}")

    def asset(self, name: str) -> str:
        return os.path.join(self.assets_dir, name)


# ---------------------------------------------------------------------------
# 形状 / 文字原语
# ---------------------------------------------------------------------------
def add_rect(slide, rect, fill: str, alpha: float | None = None, corner: float | None = None):
    x, y, w, h = rect
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if corner else MSO_SHAPE.RECTANGLE
    sp = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.shadow.inherit = False
    sp.fill.solid()
    sp.fill.fore_color.rgb = rgb(fill)
    sp.line.fill.background()
    if corner:
        short = min(w, h)
        sp.adjustments[0] = max(0.0, min(0.5, corner / short)) if short else 0.1
    if alpha is not None:
        set_alpha(sp, alpha)
    return sp


def set_alpha(shape, opacity: float) -> None:
    """Set OOXML fill opacity in the inclusive range 0..1.

    python-pptx does not expose fill transparency.  The function therefore
    writes ``a:alpha`` directly, but guards against unsupported color nodes and
    removes stale alpha children so repeated calls stay deterministic.
    """
    opacity = max(0.0, min(1.0, float(opacity)))
    solid = shape._element.spPr.find(qn("a:solidFill"))
    if solid is None:
        return
    clr = solid.find(qn("a:srgbClr"))
    if clr is None:
        return
    for old_alpha in list(clr.findall(qn("a:alpha"))):
        clr.remove(old_alpha)
    alpha = clr.makeelement(qn("a:alpha"), {"val": str(int(round(opacity * 100000)))})
    clr.append(alpha)


def _set_cjk(run, font_cjk: str) -> None:
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = rPr.makeelement(qn("a:ea"), {})
        rPr.append(ea)
    ea.set("typeface", font_cjk)


def add_text(ctx: Ctx, slide, rect, text: str, size: float, color: str, *, bold=False, align="left",
             anchor="top", min_size: float | None = None, kicker=False, italic=False, slide_no=0,
             space_after: float = 0.0) -> tuple[object, float]:
    """放一个文字框；必要时缩字。返回 (shape, 实际字号)。"""
    x, y, w, h = rect
    final = size
    overflow = False
    if min_size is not None:
        final, overflow = fit_size(text, w, h, size, min_size, ctx.ls)
        if overflow:
            ctx.warn(slide_no, f"文字放不下（已缩到 {final:.0f}pt）：「{str(text)[:30]}…」 → 删字")
        elif final < size:
            ctx.warn(slide_no, f"文字缩到 {final:.0f}pt（原 {size}pt）：「{str(text)[:30]}…」 → 考虑删字")
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tb.name = "PST_TEXT"
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = Inches(0.02)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = ANCHOR[anchor]
    paras = str(text).split("\n")
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = ALIGN[align]
        p.line_spacing = ctx.ls
        if space_after:
            p.space_after = Pt(space_after)
        run = p.add_run()
        run.text = para.upper() if (kicker and not CJK_RE.search(para)) else para
        run.font.size = Pt(final)
        run.font.bold = bold
        run.font.italic = italic
        run.font.name = ctx.fonts["font_latin"]
        run.font.color.rgb = rgb(color)
        _set_cjk(run, ctx.fonts["font_cjk"])
        if kicker:
            run._r.get_or_add_rPr().set("spc", "300")  # 字距 3pt
    return tb, final


def add_bullets(ctx: Ctx, slide, rect, items: list, size: float, color: str, accent: str, *,
                min_size: float, slide_no: int) -> object:
    x, y, w, h = rect
    joined = "\n".join(str(b) for b in items)
    final, overflow = fit_size(joined, w - 0.4, h - 0.1 * len(items), size, min_size, ctx.ls)
    if overflow:
        ctx.warn(slide_no, "bullets 放不下 → 删条或删字")
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tb.name = "PST_BULLETS"
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = Inches(0.02)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = ctx.ls
        p.space_after = Pt(final * 0.55)
        dash = p.add_run()
        dash.text = "—  "
        dash.font.size = Pt(final)
        dash.font.bold = True
        dash.font.color.rgb = rgb(accent)
        dash.font.name = ctx.fonts["font_latin"]
        run = p.add_run()
        run.text = str(item)
        run.font.size = Pt(final)
        run.font.color.rgb = rgb(color)
        run.font.name = ctx.fonts["font_latin"]
        _set_cjk(run, ctx.fonts["font_cjk"])
    return tb


def add_image_cover(ctx: Ctx, slide, rect, image_name: str, focus: str, slide_no: int):
    x, y, w, h = rect
    path = ctx.asset(image_name)
    if not os.path.isfile(path):
        ctx.warn(slide_no, f"缺图 {image_name}，已放占位块")
        add_rect(slide, rect, ctx.theme["rule"])
        add_text(ctx, slide, (x + 0.3, y + 0.3, w - 0.6, 0.6), f"缺图：{image_name}", ctx.sizes["caption"],
                 ctx.theme["text_muted"])
        return None
    fx, fy = ctx.reg["image_focus"].get(focus or "center", [0.5, 0.5])
    cropped = cover_crop(path, w, h, (fx, fy), ctx.tmp)
    with Image.open(cropped) as im:
        eff_ppi = im.size[0] / w
    if eff_ppi < 110:
        ctx.warn(slide_no, f"{image_name} 有效分辨率 {eff_ppi:.0f} ppi < 110，投影会糊")
    pic = slide.shapes.add_picture(cropped, Inches(x), Inches(y), Inches(w), Inches(h))
    pic.name = "PST_IMAGE"
    return pic


def add_scrim(ctx: Ctx, slide, rect, kind: str):
    sc = ctx.tokens["scrim"]
    if kind == "gradient":
        g = sc["gradient"]
        path = os.path.join(ctx.tmp, f"scrim_gradient_{ctx.theme[g['color']].lstrip('#')}.png")
        if not os.path.isfile(path):
            gradient_png(path, ctx.theme[g["color"]], g["alpha_from"], g["alpha_to"], g["cover_ratio"])
        x, y, w, h = rect
        picture = slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
        picture.name = "PST_SCRIM"
        return picture
    if kind in ("dark", "light"):
        s = sc[kind]
        return add_rect(slide, rect, ctx.theme[s["color"]], alpha=s["alpha"])
    return None


def add_footer(ctx: Ctx, slide, slide_no: int, total: int, colors: dict, lay: dict) -> None:
    fr = list(ctx.reg["footer"]["rect"])
    label = (ctx.deck.get("footer") or {}).get("label")
    show_num = (ctx.deck.get("footer") or {}).get("page_numbers", True)
    x, y, w, h = fr
    label_x, label_w = x, w * 0.6
    # partial-left 的左下角是图片，页脚标签移动到文字列，避免落在浅色照片上。
    if lay.get("bleed") == "partial-left":
        label_x = lay["text"][0]
        label_w = max(1.0, x + w * 0.6 - label_x)
    if lay.get("bleed") == "half-bottom":
        # Keep the footer in the text half instead of placing it over the image.
        y = min(lay["image"][1] - 0.38, lay["text"][1] + lay["text"][3] + 0.08)
    if label:
        footer_label, _ = add_text(ctx, slide, (label_x, y, label_w, h), label, ctx.sizes["footer"], colors["muted"], anchor="middle")
        footer_label.name = "PST_FOOTER_LABEL"
    if show_num:
        footer_number, _ = add_text(ctx, slide, (x + w * 0.6, y, w * 0.4, h), f"{slide_no:02d} / {total:02d}", ctx.sizes["footer"],
                                   colors["muted"], align="right", anchor="middle")
        footer_number.name = "PST_FOOTER_NUMBER"


# ---------------------------------------------------------------------------
# 版式解析
# ---------------------------------------------------------------------------
def resolve_layout(ctx: Ctx, sl: dict) -> dict:
    t, bleed = sl["type"], sl.get("bleed", "none")
    bm = ctx.reg["bleed_modes"][bleed]
    tp = ctx.reg["types"][t]
    text_rect = (tp.get("text_override") or {}).get(bleed) or bm["text"]
    scrim = sl.get("scrim") or bm.get("scrim_default") or "none"
    return {
        "type": t, "bleed": bleed, "bm": bm, "tp": tp,
        "text": tuple(text_rect),
        "image": tuple(bm["image"]) if bm.get("image") else None,
        "block": tuple(bm["block"]) if bm.get("block") else None,
        "panel": tuple(bm["panel"]) if bm.get("panel") else None,
        "on_image": bool(bm.get("on_image")),
        "scrim": scrim,
    }


def text_colors(ctx: Ctx, lay: dict) -> dict:
    th = ctx.theme
    if lay["on_image"]:
        base = th["text_on_light_scrim"] if lay["scrim"] == "light" else th["text_on_dark_scrim"]
        return {"text": base, "muted": mix(base, "#808080", 0.35), "accent": th["accent"] if lay["scrim"] != "light" else th["accent"], "rule": mix(base, "#808080", 0.6)}
    if lay["bleed"] == "color-block":
        base = th["color_block_text"]
        return {"text": base, "muted": mix(base, th["color_block"], 0.35), "accent": base, "rule": mix(base, th["color_block"], 0.55)}
    return {"text": th["text"], "muted": th["text_muted"], "accent": th["accent"], "rule": th["rule"]}


# ---------------------------------------------------------------------------
# 文字堆（label / title / body / bullets 等按顺序垂直排列，再按 anchor 定位）
# ---------------------------------------------------------------------------
def stack_text(ctx: Ctx, slide, rect, items: list[dict], anchor: str, slide_no: int) -> float:
    """items: [{text, size, color, bold, kicker, kind, min_size, gap_after}]。返回堆的总高。"""
    x, y, w, h = rect
    measured = []
    for it in items:
        if it["kind"] == "bullets":
            joined = "\n".join(str(b) for b in it["text"])
            hh = estimate_height(joined, w - 0.4, it["size"], ctx.ls) + 0.12 * len(it["text"])
        else:
            hh = estimate_height(str(it["text"]), w, it["size"], ctx.ls)
        measured.append(hh)
    total = sum(measured) + sum(it.get("gap_after", 0.0) for it in items[:-1])
    if total > h:  # 整体放不下：按比例给每项缩空间，add_text 会缩字
        scale = h / total
        measured = [m * scale for m in measured]
        total = h
    cy = {"top": y, "middle": y + (h - total) / 2, "bottom": y + h - total}[anchor]
    for it, hh in zip(items, measured):
        if it["kind"] == "bullets":
            add_bullets(ctx, slide, (x, cy, w, hh), it["text"], it["size"], it["color"], it["accent"],
                        min_size=it["min_size"], slide_no=slide_no)
        else:
            add_text(ctx, slide, (x, cy, w, hh), it["text"], it["size"], it["color"], bold=it.get("bold", False),
                     kicker=it.get("kicker", False), italic=it.get("italic", False), min_size=it.get("min_size"),
                     align=it.get("align", "left"), slide_no=slide_no)
        cy += hh + it.get("gap_after", 0.0)
    return total


def T(ctx: Ctx, text, size_key, color, **kw) -> dict:
    d = {"kind": "text", "text": text, "size": ctx.sizes[size_key], "color": color,
         "min_size": ctx.min_sizes.get(size_key)}
    d.update(kw)
    return d


# ---------------------------------------------------------------------------
# 每种 type 的构建
# ---------------------------------------------------------------------------
def build_common_background(ctx: Ctx, slide, sl: dict, lay: dict, slide_no: int) -> None:
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = rgb(ctx.theme["color_block"] if lay["bleed"] == "color-block" else ctx.theme["background"])
    if lay["image"] and sl.get("ui_mockup"):
        draw_ui_mockup(
            slide, lay["image"], sl["ui_mockup"], ctx.theme, ctx.fonts,
            instance=f"s{slide_no:02d}",
        )
    elif lay["image"] and sl.get("image"):
        add_image_cover(ctx, slide, lay["image"], sl["image"], sl.get("image_focus"), slide_no)
    if lay["on_image"] and lay["scrim"] != "none":
        add_scrim(ctx, slide, (0, 0, *ctx.canvas), lay["scrim"])
    if lay["panel"]:
        pn = ctx.tokens["scrim"]["overlay_panel"]
        add_rect(slide, lay["panel"], ctx.theme[pn["color"]], alpha=pn["alpha"], corner=pn["corner"])


def _title_band(ctx: Ctx, slide, sl: dict, lay: dict, col: dict, slide_no: int) -> tuple:
    """有内容区的类型：顶部放 label + title，返回内容区 rect。"""
    x, y, w, h = lay["text"]
    th = lay["tp"].get("title_h", 1.3)
    items = []
    if sl.get("label"):
        items.append(T(ctx, sl["label"], "label", col["accent"], kicker=True, bold=True, gap_after=0.08))
    items.append(T(ctx, sl["title"], lay["tp"]["title_size"], col["text"], bold=True))
    stack_text(ctx, slide, (x, y, w, th), items, "top", slide_no)
    return (x, y + th + 0.25, w, h - th - 0.25)


def build_text_page(ctx: Ctx, slide, sl: dict, lay: dict, col: dict, slide_no: int) -> None:
    """cover / section / statement / feature_hero / split / bullets / closing 共用。"""
    tp = lay["tp"]
    items = []
    for el in tp["elements"]:
        if el == "label" and sl.get("label"):
            items.append(T(ctx, sl["label"], "label", col["accent"], kicker=True, bold=True, gap_after=0.12))
        elif el == "title":
            items.append(T(ctx, sl["title"], tp["title_size"], col["text"], bold=True, gap_after=0.28))
        elif el == "body" and sl.get("body"):
            items.append(T(ctx, sl["body"], "body", col["muted"] if lay["type"] in ("cover", "closing") else col["text"],
                           gap_after=0.25))
        elif el == "bullets" and sl.get("bullets"):
            items.append({"kind": "bullets", "text": sl["bullets"], "size": ctx.sizes["bullet"], "color": col["text"],
                          "accent": col["accent"], "min_size": ctx.min_sizes["bullet"]})
    stack_text(ctx, slide, lay["text"], items, tp["anchor"], slide_no)
    if lay["type"] == "section":
        # 章节页：标题下一根短强调线
        x, y, w, h = lay["text"]
        add_rect(slide, (x, y + h / 2 + 0.9, 1.2, 0.06), col["accent"])


def build_quote(ctx: Ctx, slide, sl: dict, lay: dict, col: dict, slide_no: int) -> None:
    x, y, w, h = lay["text"]
    # 大引号作为装饰，用强调色
    add_text(ctx, slide, (x - 0.1, y - 0.2, 1.2, 1.2), "“", 96, col["accent"], bold=True)
    items = [
        T(ctx, sl["quote"], "quote", col["text"], bold=False, gap_after=0.3),
        T(ctx, sl.get("attribution", ""), "attribution", col["muted"], gap_after=0.5),
        T(ctx, sl["title"], "subtitle", col["muted"], bold=True),
    ]
    stack_text(ctx, slide, (x + 0.6, y + 0.6, w - 0.6, h - 0.6), items, "middle", slide_no)


def build_big_number(ctx: Ctx, slide, sl: dict, lay: dict, col: dict, slide_no: int) -> None:
    items = []
    if sl.get("label"):
        items.append(T(ctx, sl["label"], "label", col["accent"], kicker=True, bold=True, gap_after=0.1))
    items.append({"kind": "text", "text": str(sl["number"]), "size": ctx.sizes["big_number"], "color": col["accent"],
                  "bold": True, "min_size": 72, "gap_after": 0.1})
    items.append(T(ctx, sl["title"], "subtitle", col["text"], bold=True, gap_after=0.15))
    if sl.get("caption"):
        items.append(T(ctx, sl["caption"], "big_number_caption", col["muted"], gap_after=0.4))
    src = sl.get("source")
    if src:
        prefix = "Source: " if str(ctx.deck.get("language", "zh")).lower().startswith("en") else "来源："
        items.append(T(ctx, f"{prefix}{src}", "caption", col["muted"]))
    stack_text(ctx, slide, lay["text"], items, "middle", slide_no)


def build_compare(ctx: Ctx, slide, sl: dict, lay: dict, col: dict, slide_no: int) -> None:
    cx, cy, cw, ch = _title_band(ctx, slide, sl, lay, col, slide_no)
    gap = 0.8
    colw = (cw - gap) / 2
    for i, side in enumerate(("left", "right")):
        blk = sl[side]
        x = cx + i * (colw + gap)
        head_color = col["muted"] if side == "left" else col["accent"]
        add_text(ctx, slide, (x, cy, colw, 0.6), blk["title"], ctx.sizes["compare_head"], head_color, bold=True,
                 anchor="bottom", slide_no=slide_no)
        add_rect(slide, (x, cy + 0.7, colw, 0.03), head_color if side == "right" else col["rule"])
        if blk.get("bullets"):
            add_bullets(ctx, slide, (x, cy + 0.95, colw, ch - 1.0), blk["bullets"], ctx.sizes["bullet"],
                        col["text"] if side == "right" else col["muted"], head_color,
                        min_size=ctx.min_sizes["bullet"], slide_no=slide_no)


def build_flow(ctx: Ctx, slide, sl: dict, lay: dict, col: dict, slide_no: int) -> None:
    rect = _title_band(ctx, slide, sl, lay, col, slide_no)
    th = ctx.theme
    if lay["bleed"] == "color-block":
        base = th["color_block_text"]
        pal = {"node_fill": mix(th["color_block"], base, 0.1), "node_line": mix(th["color_block"], base, 0.35),
               "text": base, "muted": mix(base, th["color_block"], 0.3), "hl_fill": base,
               "hl_text": th["color_block"], "arrow": mix(th["color_block"], base, 0.5)}
    else:
        pal = {"node_fill": th["surface"], "node_line": th["rule"], "text": th["text"], "muted": th["text_muted"],
               "hl_fill": th["accent"], "hl_text": on_color(th["accent"]), "arrow": th["neutral_data"]}
    support_images = sl.get("images") or []
    flow_rect = rect
    if support_images:
        x, y, w, h = rect
        split_gap = 0.38
        flow_w = w * 0.58
        flow_rect = (x, y, flow_w, h)
        media_x = x + flow_w + split_gap
        media_w = w - flow_w - split_gap
        image_gap = 0.18
        image_h = (h - image_gap * (len(support_images) - 1)) / len(support_images)
        for idx, name in enumerate(support_images):
            yy = y + idx * (image_h + image_gap)
            pic = add_image_cover(ctx, slide, (media_x, yy, media_w, image_h), name, "center", slide_no)
            if pic is not None:
                pic.name = f"PST_FLOW_IMAGE_{idx + 1}"
    steps = []
    for step in sl["steps"]:
        item = dict(step)
        if item.get("icon"):
            item["icon"] = ctx.asset(item["icon"])
        steps.append(item)
    draw_flow(slide, flow_rect, steps, pal, ctx.fonts, ctx.sizes, ctx.reg["flow"],
              direction=sl.get("direction", "horizontal"), style=sl.get("style", "cards"))


def build_chart(ctx: Ctx, slide, sl: dict, lay: dict, col: dict, slide_no: int) -> None:
    x, y, w, h = _title_band(ctx, slide, sl, lay, col, slide_no)
    h = h - 0.15  # 给页脚留一点
    palette = build_palette(ctx.theme, ctx.tokens.get("chart"))
    out = os.path.join(ctx.tmp, f"chart_{slide_no:02d}.png")
    render_chart(
        sl["chart"], palette, out, w, h,
        fonts={**ctx.tokens["typography"], **ctx.fonts},
        cfg=ctx.tokens.get("chart"),
        language=str(ctx.deck.get("language", "zh")),
    )
    picture = slide.shapes.add_picture(out, Inches(x), Inches(y), Inches(w), Inches(h))
    picture.name = "PST_CHART"


def build_gallery(ctx: Ctx, slide, sl: dict, lay: dict, col: dict, slide_no: int) -> None:
    x, y, w, h = _title_band(ctx, slide, sl, lay, col, slide_no)
    imgs = sl["images"]
    gap = lay["tp"].get("gap", 0.25)
    n = len(imgs)
    cw = (w - gap * (n - 1)) / n
    ch = min(h - 0.2, cw * 0.75)
    for i, name in enumerate(imgs):
        add_image_cover(ctx, slide, (x + i * (cw + gap), y, cw, ch), name, "center", slide_no)
    meta_y = y + ch + 0.16
    if sl.get("caption"):
        add_text(ctx, slide, (x, meta_y, w, 0.34), sl["caption"], ctx.sizes["caption"], col["text"],
                 min_size=11, slide_no=slide_no)
        meta_y += 0.34
    if sl.get("source"):
        prefix = "Source: " if str(ctx.deck.get("language", "zh")).lower().startswith("en") else "来源："
        add_text(ctx, slide, (x, meta_y, w, 0.32), f"{prefix}{sl['source']}", ctx.sizes["flow_desc"], col["muted"],
                 min_size=12, slide_no=slide_no)


def build_media_bars(ctx: Ctx, slide, sl: dict, lay: dict, col: dict, slide_no: int) -> None:
    """Editable short bars paired one-to-one with item-specific images."""
    x, y, w, h = _title_band(ctx, slide, sl, lay, col, slide_no)
    items = sl["items"]
    ratio = max(0.45, min(0.72, float(sl.get("bar_width_ratio", 0.60))))
    n = len(items)
    gap = 0.14
    row_h = min(0.90, (h - gap * (n - 1)) / n)
    bar_w = w * ratio
    media_w = min(2.35, max(1.55, (w - bar_w) * 0.58))
    media_left = x + bar_w + 0.45
    media_right = x + w - media_w
    for i, item in enumerate(items):
        yy = y + i * (row_h + gap)
        highlighted = bool(item.get("highlight"))
        fill = col["accent"] if highlighted else ctx.theme["surface"]
        text_color = on_color(col["accent"]) if highlighted else col["text"]
        bar = add_rect(slide, (x, yy, bar_w, row_h), fill, corner=0.08)
        bar.name = f"PST_MEDIA_BAR_{i + 1}"
        add_text(ctx, slide, (x + 0.18, yy + 0.08, bar_w * 0.38, row_h - 0.16),
                 item["label"], ctx.sizes["body"], text_color, bold=True,
                 anchor="middle", min_size=14, slide_no=slide_no)
        if item.get("desc"):
            add_text(ctx, slide, (x + bar_w * 0.43, yy + 0.08, bar_w * 0.52, row_h - 0.16),
                     item["desc"], ctx.sizes["caption"],
                     text_color if highlighted else col["muted"], anchor="middle",
                     min_size=11, slide_no=slide_no)
        if sl.get("media_side", "alternating") == "alternating":
            thumb_x = media_left if i % 2 == 0 else media_right
        else:
            thumb_x = media_right
        connector = add_rect(slide, (x + bar_w, yy + row_h / 2 - 0.01,
                                     max(0.08, thumb_x - (x + bar_w)), 0.02), col["rule"])
        connector.name = f"PST_MEDIA_BAR_LINK_{i + 1}"
        pic = add_image_cover(ctx, slide, (thumb_x, yy, media_w, row_h), item["image"],
                              item.get("image_focus", "center"), slide_no)
        if pic is not None:
            pic.name = f"PST_MEDIA_BAR_IMAGE_{i + 1}"


BUILDERS = {
    "cover": build_text_page, "section": build_text_page, "statement": build_text_page,
    "feature_hero": build_text_page, "split": build_text_page, "bullets": build_text_page,
    "closing": build_text_page, "quote": build_quote, "big_number": build_big_number,
    "compare": build_compare, "flow": build_flow, "chart": build_chart, "gallery": build_gallery,
    "media_bars": build_media_bars,
}


def build_slide(ctx: Ctx, prs, sl: dict, slide_no: int, total: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    lay = resolve_layout(ctx, sl)
    col = text_colors(ctx, lay)
    build_common_background(ctx, slide, sl, lay, slide_no)
    BUILDERS[sl["type"]](ctx, slide, sl, lay, col, slide_no)
    if sl["type"] not in ctx.reg["footer"]["hide_on"]:
        add_footer(ctx, slide, slide_no, total, col, lay)
    if sl.get("notes"):
        slide.notes_slide.notes_text_frame.text = str(sl["notes"])
    ctx.manifest_slides.append({
        "no": slide_no, "type": sl["type"], "bleed": lay["bleed"], "energy": sl["energy"], "mood": sl.get("mood"),
        "title": sl["title"], "image": sl.get("image"), "scrim": lay["scrim"] if lay["on_image"] else None,
        "is_example": bool(sl.get("is_example") or (sl.get("chart") or {}).get("is_example")),
        "topics": sl.get("topics") or [],
        "ui_mockup": bool(sl.get("ui_mockup")),
    })


# ---------------------------------------------------------------------------
def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("spec")
    ap.add_argument("out")
    ap.add_argument("--no-lint", action="store_true")
    ap.add_argument("--keep-temp", action="store_true")
    args = ap.parse_args(argv)

    try:
        spec, reg, tokens = load_yaml(args.spec), load_registry(), load_tokens()
    except Exception as e:  # noqa: BLE001
        print(f"读取失败: {e}", file=sys.stderr)
        return 2

    if not args.no_lint:
        rep = lint(spec, reg, tokens, args.spec)
        rep.print()
        if rep.fails:
            print("\nlint 有 FAIL，先修 spec。", file=sys.stderr)
            return 1

    out_dir = os.path.dirname(os.path.abspath(args.out)) or "."
    os.makedirs(out_dir, exist_ok=True)
    tmp = tempfile.mkdtemp(prefix="psp_") if not args.keep_temp else os.path.join(out_dir, "_build_tmp")
    os.makedirs(tmp, exist_ok=True)

    ctx = Ctx(spec, args.spec, reg, tokens, tmp)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(ctx.canvas[0]), Inches(ctx.canvas[1])
    prs.core_properties.title = spec["deck"]["title"]
    prs.core_properties.subject = spec["deck"].get("big_idea", "")
    prs.core_properties.comments = "Generated by product-story-pptx; edit the YAML spec and rebuild."

    slides = spec["slides"]
    for i, sl in enumerate(slides, start=1):
        try:
            build_slide(ctx, prs, sl, i, len(slides))
        except Exception as e:  # noqa: BLE001
            print(f"slide {i} 构建失败 ({sl.get('type')}): {e}", file=sys.stderr)
            if not args.keep_temp:
                shutil.rmtree(tmp, ignore_errors=True)
            return 2

    prs.save(args.out)
    lineage = dict(spec["deck"].get("lineage") or {"mode": "initial", "version": "unversioned"})
    parent = lineage.get("parent_manifest")
    if parent:
        parent_path = parent if os.path.isabs(parent) else os.path.join(os.path.dirname(os.path.abspath(args.spec)), parent)
        if os.path.isfile(parent_path):
            with open(parent_path, "rb") as parent_file:
                lineage["parent_manifest_sha256"] = hashlib.sha256(parent_file.read()).hexdigest()
    with open(args.spec, "rb") as source_file:
        source_sha = hashlib.sha256(source_file.read()).hexdigest()
    mpath = os.path.splitext(args.out)[0] + ".manifest.json"
    manifest_dir = os.path.dirname(os.path.abspath(mpath))
    source_spec_abs = os.path.abspath(args.spec)
    source_spec_dir = os.path.dirname(source_spec_abs)
    try:
        output_is_in_source_project = os.path.commonpath([source_spec_abs, manifest_dir]) == source_spec_dir
    except ValueError:
        output_is_in_source_project = False
    source_spec = (
        os.path.relpath(source_spec_abs, start=manifest_dir).replace(os.sep, "/")
        if output_is_in_source_project
        else os.path.basename(source_spec_abs)
    )
    manifest = {
        "deck": spec["deck"]["title"],
        "big_idea": spec["deck"].get("big_idea"),
        "audience": spec["deck"].get("audience"),
        "desired_action": spec["deck"].get("desired_action"),
        "theme": spec["deck"]["theme"],
        "slides": ctx.manifest_slides,
        "energy_sequence": [s["energy"] for s in slides],
        "bleed_sequence": [s.get("bleed", "none") for s in slides],
        "example_data_slides": [m["no"] for m in ctx.manifest_slides if m["is_example"]],
        "warnings": ctx.warnings,
        "fonts": ctx.fonts,
        "lineage": lineage,
        "source_spec": source_spec,
        "source_spec_sha256": source_sha,
        "regression_pages": lineage.get("regression_pages") or [],
        "topic_budgets": spec["deck"].get("topic_budgets") or [],
    }
    with open(mpath, "w", encoding="utf-8") as f:
        json.dump(manifest, f, ensure_ascii=False, indent=2)

    if not args.keep_temp:
        shutil.rmtree(tmp, ignore_errors=True)
    print(f"\n→ {args.out}\n→ {mpath}")
    print(f"energy: {' '.join(str(e) for e in manifest['energy_sequence'])}")
    if ctx.warnings:
        print(f"{len(ctx.warnings)} 条 WARN，见 manifest.warnings；下一步 render_preview.py 看图。")
    return 0


if __name__ == "__main__":
    sys.exit(main())
