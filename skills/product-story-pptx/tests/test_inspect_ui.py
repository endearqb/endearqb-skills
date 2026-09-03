from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from inspect_pptx import inspect  # noqa: E402
from ui_mockup import draw_ui_mockup  # noqa: E402


PALETTE = {
    "background": "#F6F7F9",
    "surface": "#ECEFF2",
    "text": "#101820",
    "text_muted": "#66717D",
    "accent": "#00A676",
    "rule": "#CCD3D9",
}
FONTS = {"font_latin": "Inter", "font_cjk": "Noto Sans CJK SC"}


class InspectUiTest(unittest.TestCase):
    def make_presentation(self) -> Presentation:
        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
        return prs

    def test_repeated_instances_with_different_corner_warn(self):
        prs = self.make_presentation()
        for index, corner in enumerate((0.10, 0.25), start=1):
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            draw_ui_mockup(
                slide, (1.0, 1.0, 5.8, 4.8),
                {"title": "App", "badge": "界面示意", "corner_in": corner},
                PALETTE, FONTS, f"s{index:02d}",
            )
        with tempfile.NamedTemporaryFile(suffix=".pptx") as out:
            prs.save(out.name)
            report = inspect(out.name)
        self.assertTrue(any(item["code"] == "W10" and "圆角" in item["msg"] for item in report.items))

    def test_missing_divider_is_error(self):
        prs = self.make_presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        draw_ui_mockup(
            slide, (1.0, 1.0, 5.8, 4.8),
            {"title": "App", "badge": "界面示意"},
            PALETTE, FONTS, "s01",
        )
        divider = next(shape for shape in slide.shapes if shape.name.startswith("PST_UI_MOCK_DIVIDER::"))
        divider.name = "BROKEN_DIVIDER"
        with tempfile.NamedTemporaryFile(suffix=".pptx") as out:
            prs.save(out.name)
            report = inspect(out.name)
        self.assertTrue(any(item["code"] == "E09" and "DIVIDER" in item["msg"] for item in report.items))


if __name__ == "__main__":
    unittest.main()
